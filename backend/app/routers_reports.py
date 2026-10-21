"""
API Router: Kế hoạch năm (FR-03), Ánh xạ việc trường (FR-06), Thống kê & Trích xuất (FR-05, FR-07),
Quy đổi khối lượng (FR-09), Lịch sử dự đoán ML (FR-10), Hồ sơ đánh giá 1 năm giảng viên.
"""
import io
import json
import openpyxl
from datetime import datetime, date, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, desc, case

from app.db import get_db
from app.models import (
    Task, TaskAssignment, TaskType, TaskTypeScale, User, Plan, SchoolCatalogItem,
    RiskPrediction, ExtensionRequest, ProgressUpdate, Notification,
    ST_DONE, ST_DOING, ST_NEW, EXT_PENDING
)
from app.deps import current_user, manager_or_admin, lecturer_scope, can_see_lecturer, can_see_risk
from app.academic import year_range, semester_range, semester_of, academic_year_label
from app.schemas import SchoolCatalogCreate
from app.ml_client import predict_and_record_risk

router = APIRouter(prefix="/api/reports", tags=["Reports & Tools"])


# --- Kế hoạch công việc năm (FR-03, UC-05) ---
@router.get("/plans")
def list_plans(
    academic_year: str | None = None,
    lecturer_id: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user)
):
    q = db.query(Plan)
    scope = lecturer_scope(db, user)
    if scope is not None:
        q = q.filter(Plan.lecturer_id.in_(scope))
    elif lecturer_id is not None:
        q = q.filter(Plan.lecturer_id == lecturer_id)

    if academic_year:
        q = q.filter(Plan.academic_year == academic_year)

    plans = q.order_by(desc(Plan.imported_at)).all()
    res = []
    for p in plans:
        tasks = db.query(Task).filter_by(plan_id=p.id).all()
        done_cnt = sum(1 for t in tasks if t.status == ST_DONE)
        importer = db.get(User, p.imported_by) if p.imported_by else None
        res.append({
            "id": p.id,
            "lecturer_id": p.lecturer_id,
            "lecturer_name": p.lecturer.full_name if p.lecturer else "",
            "department_name": p.lecturer.department.name if p.lecturer and p.lecturer.department else "",
            "academic_year": p.academic_year,
            "file_name": p.file_name,
            "imported_by_name": importer.full_name if importer else "Hệ thống",
            "imported_at": p.imported_at,
            "note": p.note,
            "total_tasks": len(tasks),
            "completed_tasks": done_cnt,
            "tasks": [
                {
                    "id": t.id,
                    "code": t.code,
                    "title": t.title,
                    "task_type_name": t.task_type.name if t.task_type else "",
                    "priority": t.priority,
                    "original_deadline": t.original_deadline,
                    "status": t.status
                }
                for t in tasks
            ]
        })
    return res


@router.get("/plans/template")
def download_plan_template():
    """Tải file Excel mẫu Kế hoạch công việc năm (UC-05)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KeHoachNam_2026_2027"
    ws.append(["Tên công việc", "Loại công việc", "Hạn chót (YYYY-MM-DD)", "Độ ưu tiên", "Quy mô"])
    sample_rows = [
        ["Biên soạn đề cương chi tiết học phần Học máy nâng cao", "Biên soạn đề cương học phần", "2026-11-15", "Cao", "quy mô chuẩn"],
        ["Ra đề và chấm thi kết thúc học phần Trí tuệ nhân tạo HK1", "Ra đề và chấm thi kết thúc học phần", "2026-12-28", "Cao", "quy mô lớn"],
        ["Viết bài báo khoa học đăng tạp chí Scopus Q2 năm học 2026-2027", "Viết bài báo quốc tế (ISI/Scopus)", "2027-03-20", "Trung bình", "quy mô chuẩn"],
        ["Cập nhật minh chứng kiểm định chương trình đào tạo ngành CNPM", "Tự đánh giá chương trình đào tạo", "2027-04-15", "Cao", "quy mô chuẩn"],
        ["Hướng dẫn sinh viên thực hiện khóa luận tốt nghiệp đợt 1", "Hướng dẫn khóa luận tốt nghiệp", "2027-05-30", "Trung bình", "quy mô chuẩn"],
    ]
    for r in sample_rows:
        ws.append(r)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Mau_KeHoach_CongViec_Nam_2026_2027.xlsx"}
    )


@router.post("/plans/preview", dependencies=[Depends(manager_or_admin)])
async def preview_plan_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Đọc và kiểm tra cấu trúc file Excel kế hoạch năm trước khi xác nhận lưu (UC-05 Bước 3)."""
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Định dạng file không hợp lệ. Vui lòng tải lên file Excel (.xlsx).")

    ws = wb.active
    task_types = {tt.name.lower().strip(): tt.name for tt in db.query(TaskType).all()}
    default_tt_name = next(iter(task_types.values()), "Giảng dạy học phần")

    rows = []
    errors = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        title = str(row[0]).strip()
        raw_type = str(row[1]).strip() if len(row) > 1 and row[1] else default_tt_name
        matched_type = task_types.get(raw_type.lower(), default_tt_name)
        is_exact_type = raw_type.lower() in task_types

        deadline_val = row[2] if len(row) > 2 else None
        deadline_str = None
        if isinstance(deadline_val, datetime):
            deadline_str = deadline_val.date().isoformat()
        elif isinstance(deadline_val, date):
            deadline_str = deadline_val.isoformat()
        elif deadline_val:
            try:
                deadline_str = date.fromisoformat(str(deadline_val).strip()[:10]).isoformat()
            except Exception:
                errors.append(f"Dòng {idx}: Ngày hạn chót '{deadline_val}' không đúng định dạng YYYY-MM-DD")
                deadline_str = (date.today() + timedelta(days=30)).isoformat()
        else:
            deadline_str = (date.today() + timedelta(days=30)).isoformat()

        priority = str(row[3]).strip() if len(row) > 3 and row[3] else "Trung bình"
        if priority not in ("Thấp", "Trung bình", "Cao", "Khẩn cấp"):
            priority = "Trung bình"

        scale = str(row[4]).strip() if len(row) > 4 and row[4] else "quy mô chuẩn"
        if scale not in ("quy mô nhỏ", "quy mô chuẩn", "quy mô lớn"):
            scale = "quy mô chuẩn"

        rows.append({
            "row_number": idx,
            "title": title,
            "raw_task_type": raw_type,
            "matched_task_type": matched_type,
            "is_exact_type": is_exact_type,
            "deadline": deadline_str,
            "priority": priority,
            "scale": scale
        })

    return {
        "file_name": file.filename,
        "total_rows": len(rows),
        "errors": errors,
        "rows": rows
    }


@router.post("/plans/import", dependencies=[Depends(manager_or_admin)])
async def import_plan(
    lecturer_id: int,
    academic_year: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(current_user)
):
    lecturer = db.get(User, lecturer_id)
    if not lecturer:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy giảng viên")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    plan = Plan(
        lecturer_id=lecturer_id,
        academic_year=academic_year,
        file_name=file.filename,
        imported_by=user.id,
        imported_at=datetime.now()
    )

    db.add(plan)
    db.flush()

    task_types = {tt.name.lower().strip(): tt.id for tt in db.query(TaskType).all()}
    default_tt_id = next(iter(task_types.values()))
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        title = str(row[0]).strip()
        type_name = str(row[1]).lower().strip() if len(row) > 1 and row[1] else ""
        tt_id = task_types.get(type_name, default_tt_id)

        deadline_val = row[2] if len(row) > 2 else None
        if isinstance(deadline_val, datetime):
            deadline_date = deadline_val.date()
        elif isinstance(deadline_val, date):
            deadline_date = deadline_val
        elif deadline_val:
            try:
                deadline_date = date.fromisoformat(str(deadline_val).strip()[:10])
            except Exception:
                deadline_date = date.today() + timedelta(days=30)
        else:
            deadline_date = date.today() + timedelta(days=30)

        priority = str(row[3]).strip() if len(row) > 3 and row[3] else "Trung bình"
        if priority not in ("Thấp", "Trung bình", "Cao", "Khẩn cấp"):
            priority = "Trung bình"

        scale = str(row[4]).strip() if len(row) > 4 and row[4] else "quy mô chuẩn"
        if scale not in ("quy mô nhỏ", "quy mô chuẩn", "quy mô lớn"):
            scale = "quy mô chuẩn"

        task = Task(
            code=f"KH-{academic_year[:4]}-{plan.id}-{count+1:03d}",
            title=title,
            description=f"Công việc thuộc Kế hoạch năm học {academic_year} ({file.filename})",
            task_type_id=tt_id,
            scale=scale,
            original_deadline=deadline_date,
            current_deadline=deadline_date,
            priority=priority,
            status=ST_NEW,
            created_by=user.id,
            plan_id=plan.id
        )
        db.add(task)
        db.flush()

        assign = TaskAssignment(
            task_id=task.id,
            lecturer_id=lecturer_id,
            progress_pct=0,
            status=ST_NEW
        )
        db.add(assign)
        db.flush()

        # Dự đoán rủi ro ML cho đầu việc kế hoạch năm
        if deadline_date > date.today():
            predict_and_record_risk(db, task.id, assign.id, lecturer_id, trigger="assign")

        count += 1

    notif = Notification(
        user_id=lecturer_id,
        kind="new_task",
        title=f"Đã import Kế hoạch năm học {academic_year}",
        content=f"Kế hoạch công việc năm học {academic_year} gồm {count} đầu việc đã được cập nhật vào hệ thống.",
        dedupe_key=f"plan_import:{plan.id}:{lecturer_id}"
    )
    db.add(notif)

    db.commit()
    return {"message": f"Đã import thành công {count} đầu việc kế hoạch năm {academic_year}", "plan_id": plan.id, "count": count}


# --- Danh mục cấp Trường & Ánh xạ (FR-06) ---
@router.get("/school-catalog")
def get_school_catalog(db: Session = Depends(get_db)):
    items = db.query(SchoolCatalogItem).filter_by(is_active=True).all()
    res = []
    for sci in items:
        mapped_tasks = db.query(Task).filter_by(school_catalog_id=sci.id).count()
        res.append({
            "id": sci.id,
            "code": sci.code,
            "name": sci.name,
            "conversion_hours": sci.conversion_hours,
            "task_type_id": sci.task_type_id,
            "task_type_name": sci.task_type.name if sci.task_type else None,
            "note": sci.note,
            "mapped_task_count": mapped_tasks
        })
    return res


@router.post("/school-catalog", dependencies=[Depends(manager_or_admin)])
def create_school_catalog(req: SchoolCatalogCreate, db: Session = Depends(get_db)):
    existing = db.query(SchoolCatalogItem).filter_by(code=req.code).first()
    if existing:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Mã danh mục Trường '{req.code}' đã tồn tại")
    sci = SchoolCatalogItem(
        code=req.code,
        name=req.name,
        conversion_hours=req.conversion_hours,
        task_type_id=req.task_type_id,
        note=req.note,
        is_active=True
    )
    db.add(sci)
    db.commit()
    db.refresh(sci)
    return {"message": "Đã thêm danh mục công việc cấp Trường", "id": sci.id}


@router.post("/school-catalog/map-task", dependencies=[Depends(manager_or_admin)])
def map_school_task(task_id: int, school_catalog_id: int, db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if not task:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy công việc")
    sci = db.get(SchoolCatalogItem, school_catalog_id)
    if not sci:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy mục danh mục Trường")

    task.school_catalog_id = sci.id
    db.commit()
    return {"message": f"Đã ánh xạ công việc '{task.title}' vào danh mục Trường '{sci.name}' ({sci.conversion_hours} giờ chuẩn)"}


# --- Thống kê & Tính khối lượng quy đổi (FR-05, FR-07, FR-09) ---
@router.get("/statistics")
def get_statistics(
    academic_year: str | None = None,
    semester: str | None = None,
    department_id: int | None = None,
    lecturer_id: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user)
):
    query = db.query(TaskAssignment).join(Task)
    scope = lecturer_scope(db, user)
    if scope is not None:
        query = query.filter(TaskAssignment.lecturer_id.in_(scope))
    if lecturer_id is not None:
        query = query.filter(TaskAssignment.lecturer_id == lecturer_id)
    elif department_id:
        query = query.join(User, TaskAssignment.lecturer_id == User.id).filter(User.department_id == department_id)

    if academic_year:
        if semester:
            s_date, e_date = semester_range(academic_year, semester)
            query = query.filter(Task.original_deadline >= s_date.date(), Task.original_deadline < e_date.date())
        else:
            s_date, e_date = year_range(academic_year)
            query = query.filter(Task.assigned_at >= s_date, Task.assigned_at < e_date)

    assignments = query.all()
    total = len(assignments)
    done = sum(1 for a in assignments if a.status == ST_DONE)
    doing = sum(1 for a in assignments if a.status == ST_DOING)
    new_task = sum(1 for a in assignments if a.status == ST_NEW)
    late = sum(1 for a in assignments if a.status == ST_DONE and a.is_late)
    on_time = sum(1 for a in assignments if a.status == ST_DONE and not a.is_late)
    overdue_open = sum(1 for a in assignments if a.status != ST_DONE and a.task.current_deadline < date.today())
    due_soon_open = sum(
        1 for a in assignments
        if a.status != ST_DONE and 0 <= (a.task.current_deadline - date.today()).days <= 3
    )

    all_scales = db.query(TaskTypeScale).all()
    scales = {
        (tts.task_type_id, tts.scale): (tts.conversion_hours or 0.0)
        for tts in all_scales
    }
    default_scales = {}
    for tts in all_scales:
        if tts.task_type_id not in default_scales:
            default_scales[tts.task_type_id] = tts.conversion_hours or 0.0

    total_workload = 0.0
    by_group = {}
    for a in assignments:
        hrs = (
            a.task.school_item.conversion_hours
            if a.task.school_item
            else scales.get((a.task.task_type_id, a.task.scale), default_scales.get(a.task.task_type_id, 0.0))
        )
        total_workload += hrs
        grp = a.task.task_type.task_group if a.task.task_type else "Khác"
        if grp not in by_group:
            by_group[grp] = {"group": grp, "total": 0, "done": 0, "late": 0, "hours": 0.0}
        by_group[grp]["total"] += 1
        if a.status == ST_DONE:
            by_group[grp]["done"] += 1
            if a.is_late:
                by_group[grp]["late"] += 1
        by_group[grp]["hours"] = round(by_group[grp]["hours"] + hrs, 2)

    return {
        "total_tasks": total,
        "completed": done,
        "completed_on_time": on_time,
        "in_progress": doing,
        "new": new_task,
        "late_completed": late,
        "overdue_open": overdue_open,
        "due_soon_open": due_soon_open,
        "completion_rate": round(done / total * 100, 1) if total > 0 else 0.0,
        "on_time_rate": round(on_time / done * 100, 1) if done > 0 else 0.0,
        "total_converted_hours": round(total_workload, 2),
        "by_group": list(by_group.values())
    }


# --- Hồ sơ Đánh giá 1 Năm Công việc của Giảng viên (1-Year Lecturer Evaluation Profile) ---
@router.get("/lecturer-evaluation/{lecturer_id}")
def get_lecturer_evaluation(
    lecturer_id: int,
    academic_year: str = "2026-2027",
    db: Session = Depends(get_db),
    user: User = Depends(current_user)
):
    """
    Hồ sơ giả lập & đánh giá toàn diện công việc 1 năm học của Giảng viên:
    - Thông tin hồ sơ giảng viên & định mức giờ giảng dạy
    - Tổng hợp KPI hoàn thành đúng hạn / hoàn thành sớm / trễ hạn / đang thực hiện / quá hạn
    - Bảng phân rã theo Học kỳ (HK1, HK2, Học kỳ Hè) và Nhóm công việc
    - Khối lượng quy đổi giờ chuẩn (cấp Khoa + ánh xạ cấp Trường + giảng dạy)
    - Ma trận Tác vụ – Nhân sự (so sánh tỷ lệ trễ của cá nhân với trung bình Bộ môn & Khoa)
    - Bảng đối chiếu Dự đoán Rủi ro Học máy (TabPFN v2 / GBDT) vs Kết quả thực tế trên từng đầu việc
    """
    if not can_see_lecturer(db, user, lecturer_id):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Bạn không có quyền xem hồ sơ đánh giá của giảng viên này")

    lec = db.get(User, lecturer_id)
    if not lec:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Không tìm thấy giảng viên")

    prof = lec.profile
    dept = lec.department

    s_date, e_date = year_range(academic_year)
    assignments = (
        db.query(TaskAssignment)
        .join(Task)
        .filter(
            TaskAssignment.lecturer_id == lecturer_id,
            Task.assigned_at >= s_date,
            Task.assigned_at < e_date
        )
        .order_by(Task.original_deadline.asc())
        .all()
    )

    all_tts = db.query(TaskTypeScale).all()
    scales = {
        (tts.task_type_id, tts.scale): (tts.conversion_hours or 0.0, tts.expected_duration_days or 7.0)
        for tts in all_tts
    }
    default_tts = {}
    for tts in all_tts:
        if tts.task_type_id not in default_tts:
            default_tts[tts.task_type_id] = (tts.conversion_hours or 0.0, tts.expected_duration_days or 7.0)

    total = len(assignments)
    done_early = 0
    done_on_time = 0
    done_late = 0
    in_progress = 0
    new_cnt = 0
    overdue_open = 0
    due_soon_open = 0

    faculty_hours = 0.0
    school_hours = 0.0
    completed_hours = 0.0

    semesters = {
        "HK1": {"semester": "HK1 (01/09 - 15/01)", "total": 0, "done": 0, "on_time": 0, "late": 0, "open": 0, "hours": 0.0},
        "HK2": {"semester": "HK2 (16/01 - 31/05)", "total": 0, "done": 0, "on_time": 0, "late": 0, "open": 0, "hours": 0.0},
        "HE": {"semester": "Học kỳ Hè (01/06 - 31/08)", "total": 0, "done": 0, "on_time": 0, "late": 0, "open": 0, "hours": 0.0},
    }
    group_stats = {}
    type_stats = {}
    task_records = []
    ml_audit_records = []

    for a in assignments:
        t = a.task
        tt = t.task_type
        scale_hrs, exp_days = scales.get(
            (t.task_type_id, t.scale),
            default_tts.get(t.task_type_id, (10.0, 7.0))
        )
        if t.school_item:
            conv_hrs = t.school_item.conversion_hours
            school_hours += conv_hrs
        else:
            conv_hrs = scale_hrs
            faculty_hours += conv_hrs

        if a.status == ST_DONE:
            completed_hours += conv_hrs
            if a.is_late:
                done_late += 1
            elif (a.delay_days or 0) < 0:
                done_early += 1
                done_on_time += 1
            else:
                done_on_time += 1
        elif a.status == ST_DOING:
            in_progress += 1
        else:
            new_cnt += 1

        if a.status != ST_DONE:
            days_left = (t.current_deadline - date.today()).days
            if days_left < 0:
                overdue_open += 1
            elif days_left <= 5:
                due_soon_open += 1

        sem_key = semester_of(t.original_deadline)
        sem_obj = semesters.get(sem_key, semesters["HK1"])
        sem_obj["total"] += 1
        sem_obj["hours"] = round(sem_obj["hours"] + conv_hrs, 2)
        if a.status == ST_DONE:
            sem_obj["done"] += 1
            if a.is_late:
                sem_obj["late"] += 1
            else:
                sem_obj["on_time"] += 1
        else:
            sem_obj["open"] += 1

        grp = tt.task_group if tt else "Khác"
        if grp not in group_stats:
            group_stats[grp] = {"group": grp, "total": 0, "done": 0, "late": 0, "hours": 0.0}
        group_stats[grp]["total"] += 1
        group_stats[grp]["hours"] = round(group_stats[grp]["hours"] + conv_hrs, 2)
        if a.status == ST_DONE:
            group_stats[grp]["done"] += 1
            if a.is_late:
                group_stats[grp]["late"] += 1

        # Thống kê theo loại việc cho Ma trận tác vụ - nhân sự
        if tt.id not in type_stats:
            type_stats[tt.id] = {
                "task_type_id": tt.id,
                "task_type_name": tt.name,
                "task_group": tt.task_group,
                "total": 0,
                "done": 0,
                "late": 0,
                "avg_delay_days": [],
                "hours": 0.0
            }
        type_stats[tt.id]["total"] += 1
        type_stats[tt.id]["hours"] = round(type_stats[tt.id]["hours"] + conv_hrs, 2)
        if a.status == ST_DONE:
            type_stats[tt.id]["done"] += 1
            if a.is_late:
                type_stats[tt.id]["late"] += 1
            if a.delay_days is not None:
                type_stats[tt.id]["avg_delay_days"].append(a.delay_days)

        # Lấy dự đoán ML
        pred = (
            db.query(RiskPrediction)
            .filter_by(assignment_id=a.id)
            .order_by(desc(RiskPrediction.id))
            .first()
        )
        ml_data = None
        if pred and can_see_risk(db, user, lecturer_id):
            top_f = json.loads(pred.top_factors) if pred.top_factors else []
            predicted_late = pred.level in ("Cao", "Trung bình") and (pred.probability or 0) >= 0.35
            is_correct = None
            if a.status == ST_DONE and pred.probability is not None:
                if pred.level == "Cao":
                    is_correct = bool(a.is_late is True)
                elif pred.level == "Thấp":
                    is_correct = bool(a.is_late is False)
                else:
                    is_correct = bool((pred.probability >= 0.33) == bool(a.is_late))

            ml_data = {
                "id": pred.id,
                "probability": pred.probability,
                "level": pred.level,
                "uncertainty": pred.uncertainty,
                "top_factors": top_f,
                "recommendation": pred.recommendation,
                "low_confidence": pred.low_confidence,
                "trigger": pred.trigger,
                "is_correct": is_correct
            }
            ml_audit_records.append({
                "task_id": t.id,
                "task_code": t.code,
                "task_title": t.title,
                "task_type_name": tt.name,
                "semester": sem_key,
                "assigned_at": t.assigned_at.date().isoformat(),
                "original_deadline": t.original_deadline.isoformat(),
                "current_deadline": t.current_deadline.isoformat(),
                "status": a.status,
                "completed_at": a.completed_at.date().isoformat() if a.completed_at else None,
                "actual_is_late": a.is_late if a.status == ST_DONE else None,
                "delay_days": a.delay_days,
                "probability": pred.probability,
                "level": pred.level,
                "uncertainty": pred.uncertainty,
                "top_factors": top_f,
                "recommendation": pred.recommendation,
                "is_correct": is_correct
            })

        task_records.append({
            "task_id": t.id,
            "code": t.code,
            "title": t.title,
            "task_type_name": tt.name,
            "task_group": tt.task_group,
            "scale": t.scale,
            "semester": sem_key,
            "priority": t.priority,
            "complexity": t.complexity,
            "assigned_at": t.assigned_at.date().isoformat(),
            "original_deadline": t.original_deadline.isoformat(),
            "current_deadline": t.current_deadline.isoformat(),
            "completed_at": a.completed_at.date().isoformat() if a.completed_at else None,
            "progress_pct": a.progress_pct,
            "status": a.status,
            "is_late": a.is_late,
            "delay_days": a.delay_days,
            "conversion_hours": round(conv_hrs, 2),
            "school_catalog_code": t.school_item.code if t.school_item else None,
            "school_catalog_name": t.school_item.name if t.school_item else None,
            "extension_count": len(t.extensions),
            "ml_prediction": ml_data
        })

    # Tính Ma trận tác vụ - nhân sự (so sánh với Bộ môn và Toàn Khoa)
    late_expr = case((TaskAssignment.is_late == True, 1), else_=0)
    matrix_rows = []
    for tt_id, st in type_stats.items():
        # Toàn Khoa
        g_q = (
            db.query(func.count(TaskAssignment.id).label("cnt"), func.sum(late_expr).label("late"))
            .join(Task)
            .filter(Task.task_type_id == tt_id, TaskAssignment.status == ST_DONE)
            .first()
        )
        g_cnt = int(g_q.cnt or 0) if g_q else 0
        g_late = int(g_q.late or 0) if g_q else 0
        glob_rate = round(g_late / g_cnt * 100, 1) if g_cnt > 0 else 18.0

        # Bộ môn
        d_rate = glob_rate
        if dept:
            d_q = (
                db.query(func.count(TaskAssignment.id).label("cnt"), func.sum(late_expr).label("late"))
                .join(Task)
                .join(User, TaskAssignment.lecturer_id == User.id)
                .filter(User.department_id == dept.id, Task.task_type_id == tt_id, TaskAssignment.status == ST_DONE)
                .first()
            )
            d_cnt = int(d_q.cnt or 0) if d_q else 0
            d_late = int(d_q.late or 0) if d_q else 0
            if d_cnt > 0:
                d_rate = round(d_late / d_cnt * 100, 1)

        lec_rate = round(st["late"] / st["done"] * 100, 1) if st["done"] > 0 else 0.0
        avg_delay = round(sum(st["avg_delay_days"]) / len(st["avg_delay_days"]), 1) if st["avg_delay_days"] else 0.0
        matrix_rows.append({
            "task_type_id": tt_id,
            "task_type_name": st["task_type_name"],
            "task_group": st["task_group"],
            "total_assigned": st["total"],
            "completed_count": st["done"],
            "late_count": st["late"],
            "lecturer_late_rate": lec_rate,
            "dept_late_rate": d_rate,
            "global_late_rate": glob_rate,
            "avg_delay_days": avg_delay,
            "converted_hours": st["hours"]
        })

    done_total = done_on_time + done_late
    teaching_hours_year = round((prof.teaching_hours_semester if prof and prof.teaching_hours_semester else 135.0) * 2, 1)
    total_task_hours = round(faculty_hours + school_hours, 2)

    # Tổng kết độ chính xác ML trên giảng viên này
    eval_preds = [r for r in ml_audit_records if r["actual_is_late"] is not None and r["probability"] is not None]
    high_preds = [r for r in eval_preds if r["level"] == "Cao"]
    high_correct = sum(1 for r in high_preds if r["actual_is_late"] is True)
    low_preds = [r for r in eval_preds if r["level"] == "Thấp"]
    low_correct = sum(1 for r in low_preds if r["actual_is_late"] is False)
    overall_correct = sum(1 for r in eval_preds if r["is_correct"] is True)

    return {
        "academic_year": academic_year,
        "lecturer": {
            "id": lec.id,
            "username": lec.username,
            "full_name": lec.full_name,
            "email": lec.email,
            "lecturer_code": prof.lecturer_code if prof else f"GV{lec.id:04d}",
            "department_name": dept.name if dept else "Khoa CNTT",
            "academic_rank": prof.academic_rank if prof else "Giảng viên",
            "degree": prof.degree if prof else "Thạc sĩ",
            "management_role": prof.management_role if prof else "Không",
            "career_start_date": prof.career_start_date.isoformat() if prof and prof.career_start_date else "2018-09-01",
            "join_date": prof.join_date.isoformat() if prof and prof.join_date else "2020-02-01",
            "teaching_hours_semester": prof.teaching_hours_semester if prof else 135.0,
            "teaching_hours_year": teaching_hours_year,
            "course_count": prof.course_count if prof else 3,
            "student_count": prof.student_count if prof else 240,
            "supervised_student_count": prof.supervised_student_count if prof else 4,
        },
        "summary": {
            "total_tasks": total,
            "completed_total": done_total,
            "completed_early": done_early,
            "completed_on_time": done_on_time,
            "completed_late": done_late,
            "in_progress": in_progress,
            "new_tasks": new_cnt,
            "overdue_open": overdue_open,
            "due_soon_open": due_soon_open,
            "completion_rate": round(done_total / total * 100, 1) if total > 0 else 0.0,
            "on_time_rate": round(done_on_time / done_total * 100, 1) if done_total > 0 else 0.0,
            "late_rate": round(done_late / done_total * 100, 1) if done_total > 0 else 0.0,
            "faculty_converted_hours": round(faculty_hours, 2),
            "school_mapped_hours": round(school_hours, 2),
            "total_task_converted_hours": total_task_hours,
            "completed_converted_hours": round(completed_hours, 2),
            "target_task_hours": 180.0,
            "workload_fulfillment_pct": round(total_task_hours / 180.0 * 100, 1)
        },
        "semesters": list(semesters.values()),
        "by_group": list(group_stats.values()),
        "task_personnel_matrix": matrix_rows,
        "ml_audit": {
            "total_evaluated": len(eval_preds),
            "overall_accuracy": round(overall_correct / len(eval_preds) * 100, 1) if eval_preds else 0.0,
            "high_risk_total": len(high_preds),
            "high_risk_actual_late": high_correct,
            "high_risk_precision": round(high_correct / len(high_preds) * 100, 1) if high_preds else 0.0,
            "low_risk_total": len(low_preds),
            "low_risk_actual_on_time": low_correct,
            "low_risk_accuracy": round(low_correct / len(low_preds) * 100, 1) if low_preds else 0.0,
            "records": ml_audit_records
        },
        "tasks": task_records
    }


# --- Trích xuất báo cáo Excel (FR-05) ---
@router.get("/export-excel")
def export_excel(
    academic_year: str = "2026-2027",
    semester: str | None = None,
    lecturer_id: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user)
):
    query = db.query(TaskAssignment).join(Task)
    scope = lecturer_scope(db, user)
    if scope is not None:
        query = query.filter(TaskAssignment.lecturer_id.in_(scope))
    elif lecturer_id is not None:
        query = query.filter(TaskAssignment.lecturer_id == lecturer_id)

    if semester:
        s_date, e_date = semester_range(academic_year, semester)
    else:
        s_date, e_date = year_range(academic_year)
    query = query.filter(Task.assigned_at >= s_date, Task.assigned_at < e_date)

    assignments = query.order_by(Task.assigned_at.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo công việc"

    headers = [
        "Mã CV", "Tên công việc", "Loại công việc", "Nhóm việc", "Danh mục Trường",
        "Giảng viên", "Mã GV", "Bộ môn", "Học kỳ", "Ngày giao", "Hạn gốc", "Hạn hiện tại",
        "Tiến độ (%)", "Trạng thái", "Trễ hạn gốc", "Độ lệch (ngày)", "Giờ chuẩn quy đổi"
    ]
    ws.append(headers)

    scales = {
        (tts.task_type_id, tts.scale): (tts.conversion_hours or 0.0)
        for tts in db.query(TaskTypeScale).all()
    }

    for a in assignments:
        conv_hrs = (
            a.task.school_item.conversion_hours
            if a.task.school_item
            else scales.get((a.task.task_type_id, a.task.scale), 0.0)
        )
        ws.append([
            a.task.code,
            a.task.title,
            a.task.task_type.name,
            a.task.task_type.task_group,
            f"{a.task.school_item.code} - {a.task.school_item.name}" if a.task.school_item else "Cấp Khoa/Bộ môn",
            a.lecturer.full_name,
            a.lecturer.profile.lecturer_code if a.lecturer.profile else "",
            a.lecturer.department.name if a.lecturer.department else "",
            semester_of(a.task.assigned_at.date()),
            str(a.task.assigned_at.date()),
            str(a.task.original_deadline),
            str(a.task.current_deadline),
            a.progress_pct,
            a.status,
            "Trễ hạn" if a.is_late else ("Đúng hạn" if a.status == ST_DONE else "Đang mở"),
            a.delay_days if a.delay_days is not None else 0,
            round(conv_hrs, 2)
        ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"BaoCao_CongViec_{academic_year}_{semester or 'CaNam'}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# --- Đối chiếu Lịch sử Dự đoán ML (FR-10) ---
@router.get("/ml-accuracy")
def get_ml_accuracy(db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Đối chiếu độ chính xác của mô hình ML giữa dự đoán và kết quả thực tế."""
    query = db.query(RiskPrediction).filter(RiskPrediction.actual_is_late != None, RiskPrediction.probability != None)

    scope = lecturer_scope(db, user)
    if scope is not None:
        query = query.filter(RiskPrediction.lecturer_id.in_(scope))

    preds = query.order_by(desc(RiskPrediction.predicted_at)).all()
    if not preds:
        return {
            "total_evaluated": 0,
            "high_risk": {"total_predicted": 0, "actual_late": 0, "precision": 0.0},
            "medium_risk": {"total_predicted": 0, "actual_late": 0, "late_rate": 0.0},
            "low_risk": {"total_predicted": 0, "actual_on_time": 0, "accuracy": 0.0},
            "recent_comparisons": []
        }

    total = len(preds)
    correct_high = sum(1 for p in preds if p.level == "Cao" and p.actual_is_late is True)
    total_high = sum(1 for p in preds if p.level == "Cao")

    late_med = sum(1 for p in preds if p.level == "Trung bình" and p.actual_is_late is True)
    total_med = sum(1 for p in preds if p.level == "Trung bình")

    correct_low = sum(1 for p in preds if p.level == "Thấp" and p.actual_is_late is False)
    total_low = sum(1 for p in preds if p.level == "Thấp")

    recent = []
    for p in preds[:15]:
        t = db.get(Task, p.task_id)
        lec = db.get(User, p.lecturer_id)
        recent.append({
            "id": p.id,
            "task_code": t.code if t else "",
            "task_title": t.title if t else "",
            "lecturer_name": lec.full_name if lec else "",
            "probability": p.probability,
            "level": p.level,
            "uncertainty": p.uncertainty,
            "actual_is_late": p.actual_is_late,
            "predicted_at": p.predicted_at
        })

    return {
        "total_evaluated": total,
        "high_risk": {
            "total_predicted": total_high,
            "actual_late": correct_high,
            "precision": round(correct_high / total_high * 100, 1) if total_high > 0 else 0.0
        },
        "medium_risk": {
            "total_predicted": total_med,
            "actual_late": late_med,
            "late_rate": round(late_med / total_med * 100, 1) if total_med > 0 else 0.0
        },
        "low_risk": {
            "total_predicted": total_low,
            "actual_on_time": correct_low,
            "accuracy": round(correct_low / total_low * 100, 1) if total_low > 0 else 0.0
        },
        "recent_comparisons": recent
    }


